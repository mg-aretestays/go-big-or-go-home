"""Build the 2026 'Go Big or Go Home' league workbook from the Sleeper API.

Usage:  python build_gbgh_2026.py [--weeks N]
  --weeks N   only include weeks 1..N (default: every week Sleeper has finished scoring)

Rerun any time after Monday night to refresh the workbook.
"""
import argparse
import io
import json
import re
import time
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LEAGUE_ID = "1312054916033835008"
SEASON = "2026"
# Export URL of the original tracker; its FF Payments tab is copied into the xlsx.
# Kept in an untracked local file because anyone with the link can open the sheet.
SOURCE_SHEET_FILE = Path(__file__).with_name(".source_sheet_url")
OUT = Path(__file__).with_name(f"Go Big or Go Home {SEASON}.xlsx")
FP_CACHE = Path(__file__).with_name(".fp_cache") / SEASON
FP_POSITIONS = ("QB", "RB", "WR", "TE")
FP_CRAWL_DELAY = 5  # seconds, per fantasypros.com/robots.txt
FP_REFRESH_WEEKS = 2  # re-download the most recent weeks to pick up stat corrections

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
ME_FILL = PatternFill("solid", fgColor="FEF3C7")
WIN_FONT = Font(bold=True, color="15803D")
LOSS_FONT = Font(bold=True, color="B91C1C")
PAID_FILL = PatternFill("solid", fgColor="DCFCE7")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
ME = "855351711852781568"  # MGardner6


def api(path):
    with urllib.request.urlopen(f"https://api.sleeper.app/v1/{path}") as r:
        return json.load(r)


def comp_rank(values):
    """1 = highest. Ties share the better rank."""
    return [1 + sum(o > v for o in values) for v in values]


def load_season(weeks_override):
    league = api(f"league/{LEAGUE_ID}")
    users = {u["user_id"]: u for u in api(f"league/{LEAGUE_ID}/users")}
    rosters = api(f"league/{LEAGUE_ID}/rosters")
    reg_weeks = league["settings"]["playoff_week_start"] - 1
    last = weeks_override or league["settings"]["last_scored_leg"]
    weeks = list(range(1, min(last, reg_weeks) + 1))

    teams = {}
    for r in rosters:
        u = users.get(r["owner_id"], {})
        name = u.get("display_name", f"Roster {r['roster_id']}")
        teams[r["roster_id"]] = {
            "rid": r["roster_id"],
            "user_id": r["owner_id"],
            "handle": f"@{name}",
            "display": name,
            "team_name": ((u.get("metadata") or {}).get("team_name") or f"Team {name}").strip(),
            "games": {},
        }

    for w in weeks:
        rows = [m for m in api(f"league/{LEAGUE_ID}/matchups/{w}") if m.get("matchup_id")]
        pts = {m["roster_id"]: round(m["points"] or 0, 2) for m in rows}
        ids = list(pts)
        ranks = dict(zip(ids, comp_rank([pts[i] for i in ids])))
        by_match = {}
        for m in rows:
            by_match.setdefault(m["matchup_id"], []).append(m["roster_id"])
        for a, b in (pair for pair in by_match.values() if len(pair) == 2):
            for me, opp in ((a, b), (b, a)):
                res = "Win" if pts[me] > pts[opp] else "Loss" if pts[me] < pts[opp] else "Tie"
                teams[me]["games"][w] = {
                    "pts": pts[me], "opp": opp, "opp_pts": pts[opp], "result": res,
                    "rank": ranks[me], "opp_rank": ranks[opp],
                    "allplay_w": sum(pts[me] > pts[o] for o in ids if o != me),
                    "allplay_l": sum(pts[me] < pts[o] for o in ids if o != me),
                }

    n = len(teams)
    for t in teams.values():
        g = t["games"].values()
        t["w"] = sum(x["result"] == "Win" for x in g)
        t["l"] = sum(x["result"] == "Loss" for x in g)
        t["t"] = sum(x["result"] == "Tie" for x in g)
        t["pf"] = round(sum(x["pts"] for x in g), 2)
        t["pa"] = round(sum(x["opp_pts"] for x in g), 2)
        gp = max(len(t["games"]), 1)
        t["pf_avg"] = round(t["pf"] / gp, 2)
        t["pa_avg"] = round(t["pa"] / gp, 2)
        t["diff_avg"] = round(t["pf_avg"] - t["pa_avg"], 2)
        t["ap_w"] = sum(x["allplay_w"] for x in g)
        t["ap_l"] = sum(x["allplay_l"] for x in g)
        ap_games = t["ap_w"] + t["ap_l"]
        t["exp_w"] = round(t["ap_w"] / ap_games * len(t["games"]), 2) if ap_games else 0
        t["top_half"] = sum(x["rank"] <= n // 2 for x in g)
        # Same luck definition as the 2025 sheet: +1 for a win while scoring in the
        # bottom half, -1 for a loss while scoring in the top half.
        t["luck_events"] = {}
        for wk, x in t["games"].items():
            if x["result"] == "Win" and x["rank"] > n // 2:
                t["luck_events"][wk] = f"WIN v{x['opp_rank']}"
            elif x["result"] == "Loss" and x["rank"] <= n // 2:
                t["luck_events"][wk] = f"LOST v{x['opp_rank']}"
        t["luck"] = sum(1 if v.startswith("WIN") else -1 for v in t["luck_events"].values())

    order = sorted(teams.values(), key=lambda t: (-t["w"], -t["t"], -t["pf"]))
    for i, t in enumerate(order, 1):
        t["place"] = i
    for key, rk in (("pf", "pf_rank"), ("pa", "pa_rank")):
        for t, r in zip(order, comp_rank([t[key] for t in order])):
            t[rk] = r
    return league, weeks, reg_weeks, order, teams


def header(ws, row, col, values):
    for i, v in enumerate(values):
        c = ws.cell(row, col + i, v)
        c.fill, c.font, c.alignment, c.border = HEADER_FILL, HEADER_FONT, CENTER, BOX


def title(ws, text):
    ws["B1"] = text
    ws["B1"].font = TITLE_FONT


def heat_rule():
    return ColorScaleRule(start_type="num", start_value=1, start_color="22C55E",
                          mid_type="num", mid_value=6.5, mid_color="FDE68A",
                          end_type="num", end_value=12, end_color="EF4444")


def build_heatmap(wb, weeks, reg_weeks, order):
    ws = wb.active
    ws.title = "HeatMap"
    title(ws, f"Go Big or Go Home – {SEASON} Heat Map (weekly scoring rank, 1 = top score)")
    n = len(order)
    header(ws, 2, 2, ["Place"] + list(range(1, n + 1)))
    header(ws, 3, 2, ["Team"] + [t["handle"] for t in order])
    luck_col = n + 4
    header(ws, 2, luck_col, ["Luck notes"] + [""] * n)
    header(ws, 3, luck_col, ["Team"] + [t["handle"] for t in order])

    for wk in range(1, reg_weeks + 1):
        r = 3 + wk
        ws.cell(r, 2, f"Week {wk:02d}").font = Font(bold=True)
        ws.cell(r, luck_col, f"Week {wk:02d}").font = Font(bold=True)
        for i, t in enumerate(order):
            g = t["games"].get(wk)
            c = ws.cell(r, 3 + i, g["rank"] if g else None)
            c.alignment, c.border = CENTER, BOX
            note = t["luck_events"].get(wk)
            if note:
                lc = ws.cell(r, luck_col + 1 + i, note)
                lc.font = WIN_FONT if note.startswith("WIN") else LOSS_FONT
                lc.alignment = CENTER
                c.font = Font(bold=True, underline="single")
    last_row = 3 + reg_weeks
    ws.conditional_formatting.add(
        f"C4:{get_column_letter(2 + n)}{last_row}", heat_rule())

    r = last_row + 2
    summary = [
        ("Record", lambda t: f"{t['w']}-{t['l']}" + (f"-{t['t']}" if t["t"] else "")),
        ("Record v Everyone", lambda t: f"{t['ap_w']}-{t['ap_l']}"),
        ("Expected Wins (all-play)", lambda t: t["exp_w"]),
        ("Wins Above Expected", lambda t: round(t["w"] - t["exp_w"], 2)),
        ("Luck (sheet method)", lambda t: t["luck"]),
        ("Avg. Points For", lambda t: t["pf_avg"]),
        ("Points For Rank", lambda t: t["pf_rank"]),
        ("Avg. Points Against", lambda t: t["pa_avg"]),
        ("Points Against Rank (1 = toughest)", lambda t: t["pa_rank"]),
        ("Avg. +/- Points Diff.", lambda t: t["diff_avg"]),
        ("Weeks in Top 1/2 of Scoring", lambda t: t["top_half"]),
    ]
    for label, fn in summary:
        ws.cell(r, 2, label).font = Font(bold=True)
        for i, t in enumerate(order):
            c = ws.cell(r, 3 + i, fn(t))
            c.alignment, c.border = CENTER, BOX
        r += 1

    ws.cell(r + 1, 2, "Luck notes: a WIN while scoring in the bottom half (+1) or a LOSS while "
                      "scoring in the top half (-1). 'v3' = opponent's scoring rank that week.").font = Font(italic=True, color="6B7280")
    ws.column_dimensions["B"].width = 34
    for col in range(3, luck_col + n + 2):
        ws.column_dimensions[get_column_letter(col)].width = 15
    ws.column_dimensions[get_column_letter(luck_col)].width = 12
    ws.freeze_panes = "C4"


def build_weekly(wb, weeks, reg_weeks, order, teams):
    ws = wb.create_sheet("WeeklyScores")
    title(ws, f"Weekly Scores – {SEASON}")
    header(ws, 3, 2, ["Team", "Week", "Win / Loss", "Match Score", "Score", "Rank",
                      "Opponent", "Opponent Score", "Opp. Rank"])
    r = 4
    for t in order:
        for wk in range(1, reg_weeks + 1):
            g = t["games"].get(wk)
            row = [t["handle"], f"{wk:02d}"]
            if g:
                row += [g["result"], f"{g['pts']:.2f} - {g['opp_pts']:.2f}", g["pts"], g["rank"],
                        teams[g["opp"]]["handle"], g["opp_pts"], g["opp_rank"]]
            for i, v in enumerate(row):
                c = ws.cell(r, 2 + i, v)
                c.border = BOX
                if i == 2 and v in ("Win", "Loss"):
                    c.font = WIN_FONT if v == "Win" else LOSS_FONT
            r += 1
        ws.cell(r, 2, f"{t['handle']} total").font = Font(bold=True)
        ws.cell(r, 4, f"{t['w']}-{t['l']}").font = Font(bold=True)
        ws.cell(r, 6, t["pf"]).font = Font(bold=True)
        ws.cell(r, 9, t["pa"]).font = Font(bold=True)
        r += 2

    sc = 13
    header(ws, 3, sc, ["Place", "Team", "Team Name", "Wins", "Losses", "Points For", "PF Avg",
                       "PF Rank", "Points Against", "PA Avg", "PA Rank", "+/- Avg"])
    for i, t in enumerate(order):
        vals = [t["place"], t["handle"], t["team_name"], t["w"], t["l"], t["pf"], t["pf_avg"],
                t["pf_rank"], t["pa"], t["pa_avg"], t["pa_rank"], t["diff_avg"]]
        for j, v in enumerate(vals):
            c = ws.cell(4 + i, sc + j, v)
            c.border = BOX
            if t["user_id"] == ME:
                c.fill = ME_FILL
    widths = {2: 18, 3: 7, 4: 11, 5: 17, 6: 9, 7: 7, 8: 18, 9: 15, 10: 10,
              sc: 7, sc + 1: 18, sc + 2: 28}
    for col, wdt in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = wdt
    for col in range(sc + 3, sc + 12):
        ws.column_dimensions[get_column_letter(col)].width = 13
    ws.freeze_panes = "B4"


def build_pfpa(wb, weeks, reg_weeks, order):
    ws = wb.create_sheet("PFPA")
    title(ws, "POINTS FOR")
    cols = [f"{w:02d}" for w in range(1, reg_weeks + 1)]

    def grid(start_row, key):
        header(ws, start_row, 2, ["Team"] + cols + ["Total", "Average"])
        for i, t in enumerate(order):
            r = start_row + 1 + i
            ws.cell(r, 2, t["handle"]).border = BOX
            vals = [t["games"][w][key] if w in t["games"] else None for w in range(1, reg_weeks + 1)]
            for j, v in enumerate(vals):
                c = ws.cell(r, 3 + j, v)
                c.border, c.number_format = BOX, "0.00"
            played = [v for v in vals if v is not None]
            ws.cell(r, 3 + reg_weeks, round(sum(played), 2)).font = Font(bold=True)
            ws.cell(r, 4 + reg_weeks, round(sum(played) / len(played), 2) if played else None).font = Font(bold=True)
        return start_row + len(order)

    end = grid(2, "pts")
    ws.conditional_formatting.add(
        f"C3:{get_column_letter(2 + reg_weeks)}{end}",
        ColorScaleRule(start_type="min", start_color="EF4444", mid_type="percentile",
                       mid_value=50, mid_color="FDE68A", end_type="max", end_color="22C55E"))
    pa_title = end + 3
    ws.cell(pa_title, 2, "POINTS AGAINST").font = TITLE_FONT
    grid(pa_title + 1, "opp_pts")
    ws.column_dimensions["B"].width = 18
    for col in range(3, reg_weeks + 5):
        ws.column_dimensions[get_column_letter(col)].width = 9


def build_league_data(wb, weeks, order, teams):
    ws = wb.create_sheet("League Data")
    header(ws, 1, 1, ["Week", "Team", "Team Name", "Win / Loss", "Score", "Opponent",
                      "Opponent Score", "Match Score", "Margin", "Weekly Rank"])
    r = 2
    for wk in weeks:
        for t in sorted(order, key=lambda t: t["games"].get(wk, {}).get("rank", 99)):
            g = t["games"].get(wk)
            if not g:
                continue
            vals = [f"{wk:02d}", t["handle"], t["team_name"], g["result"], g["pts"],
                    teams[g["opp"]]["handle"], g["opp_pts"],
                    f"{g['pts']:.2f} - {g['opp_pts']:.2f}", round(g["pts"] - g["opp_pts"], 2), g["rank"]]
            for j, v in enumerate(vals):
                ws.cell(r, 1 + j, v).border = BOX
            r += 1
    for col, wdt in zip(range(1, 11), (7, 18, 28, 11, 9, 18, 15, 17, 9, 12)):
        ws.column_dimensions[get_column_letter(col)].width = wdt
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{r - 1}"


def build_managers(wb, order):
    ws = wb.create_sheet("Managers")
    title(ws, "League Managers (from Sleeper)")
    header(ws, 3, 2, ["Place", "Sleeper Handle", "Team Name", "Roster ID", "Sleeper User ID"])
    for i, t in enumerate(order):
        for j, v in enumerate([t["place"], t["handle"], t["team_name"], t["rid"], t["user_id"]]):
            ws.cell(4 + i, 2 + j, v).border = BOX
    ws.cell(6 + len(order), 2, "@HK141 was @HarrisonKakkuri in 2025 (same Sleeper user ID).").font = Font(italic=True, color="6B7280")
    for col, wdt in zip(range(2, 7), (7, 18, 28, 10, 22)):
        ws.column_dimensions[get_column_letter(col)].width = wdt


def build_payments(wb, order):
    """Copy the dues tracker from the original sheet, keyed to current Sleeper handles."""
    if not SOURCE_SHEET_FILE.exists():
        print(f"FF Payments skipped: no {SOURCE_SHEET_FILE.name}")
        return
    try:
        with urllib.request.urlopen(SOURCE_SHEET_FILE.read_text().strip()) as r:
            src = load_workbook(io.BytesIO(r.read()), data_only=True)["FF Payments"]
    except Exception as e:  # sheet unshared or offline: skip rather than fail the build
        print(f"FF Payments not copied: {e}")
        return
    renames = {"@HarrisonKakkuri": "@HK141"}
    ws = wb.create_sheet("FF Payments")
    title(ws, "Go Big or Go Home Yearly Payments")
    rows = list(src.iter_rows(values_only=True))
    start = next(i for i, row in enumerate(rows) if "Season" in row)
    for out_r, row in enumerate((r for r in rows[start:] if any(v is not None for v in r)), 3):
        for j, v in enumerate(row):
            if isinstance(v, float) and v.is_integer():
                v = int(v)
            v = renames.get(v, v)
            c = ws.cell(out_r, 1 + j, v)
            if out_r == 3 and v is not None:
                c.fill, c.font, c.alignment = HEADER_FILL, HEADER_FONT, CENTER
            elif v is True:
                c.fill = PAID_FILL
            c.border = BOX if v is not None else Border()
    for col in range(2, 18):
        ws.column_dimensions[get_column_letter(col)].width = 15


class _StatsTable(HTMLParser):
    """Pulls the #data table off a FantasyPros stats page."""

    def __init__(self):
        super().__init__()
        self.in_table = False
        self.groups, self.cols, self.rows = [], [], []
        self.section = None  # 'tier' | 'head' | 'body'
        self.cell, self.row, self.fp_id, self.colspan = None, None, None, 1

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag == "table" and a.get("id") == "data":
            self.in_table = True
        if not self.in_table:
            return
        if tag == "tr":
            cls = a.get("class") or ""
            self.section = "tier" if "tier-row" in cls else ("body" if "mpb-player" in cls else "head")
            self.fp_id = cls.replace("mpb-player-", "") if self.section == "body" else None
            self.row = []
        elif tag in ("td", "th"):
            self.cell, self.colspan = "", int(a.get("colspan", 1))

    def handle_data(self, data):
        if self.in_table and self.cell is not None:
            self.cell += data

    def handle_endtag(self, tag):
        if not self.in_table:
            return
        if tag in ("td", "th") and self.cell is not None:
            text = " ".join(self.cell.split())
            if self.section == "tier":
                self.groups += [text] * self.colspan
            else:
                self.row.append(text)
            self.cell = None
        elif tag == "tr" and self.row is not None:
            if self.section == "head" and self.row:
                self.cols = self.row
            elif self.section == "body":
                self.rows.append((self.fp_id, self.row))
            self.row = None
        elif tag == "table":
            self.in_table = False


def fp_page(pos, week, last_week):
    """Return the HTML for one position/week, using the local cache for settled weeks."""
    FP_CACHE.mkdir(parents=True, exist_ok=True)
    path = FP_CACHE / f"{pos.lower()}_w{week:02d}.html"
    if path.exists() and week <= last_week - FP_REFRESH_WEEKS:
        return path.read_text(encoding="utf-8")
    time.sleep(FP_CRAWL_DELAY)
    url = (f"https://www.fantasypros.com/nfl/stats/{pos.lower()}.php"
           f"?year={SEASON}&range=week&week={week}&scoring=HALF")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 (league stats sheet)"})
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                html = r.read().decode("utf-8")
            break
        except OSError:
            if attempt == 2:
                raise
            time.sleep(FP_CRAWL_DELAY * (attempt + 2))
    if "Half PPR" not in html[:2000]:
        raise RuntimeError(f"{url} did not return Half PPR scoring")
    path.write_text(html, encoding="utf-8")
    return html


def _num(v):
    v = v.replace(",", "").rstrip("%")
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except ValueError:
        return v


def load_fantasypros(weeks):
    """{pos: {'cols': [(group, name)...], 'rows': [dict...]}} of players who played each week."""
    out = {}
    for pos in FP_POSITIONS:
        cols, rows = None, []
        for wk in weeks:
            p = _StatsTable()
            p.feed(fp_page(pos, wk, weeks[-1]))
            groups = [""] * 2 + p.groups[2:]  # Rank, Player have no group
            names = list(zip(groups, p.cols))
            keep = [i for i, (_, n) in enumerate(names) if i >= 2 and n not in ("FPTS/G", "ROST")]
            cols = [names[i] for i in keep]
            for fp_id, cells in p.rows:
                m = re.match(r"(.+?)\s*\((\w+)\)$", cells[1])
                player, team = (m.group(1), m.group(2)) if m else (cells[1], "FA")
                stats = {names[i]: _num(cells[i]) for i in keep}
                if not stats.get(("MISC", "G")):
                    continue
                rows.append({"fp_id": fp_id, "week": wk, "player": player, "team": team, "stats": stats})
        out[pos] = {"cols": cols, "rows": rows}
    return out


# FantasyPros name -> Sleeper name, where they differ beyond punctuation/suffixes.
FP_NAME_ALIASES = {"hollywood brown": "marquise brown"}


def _norm(name):
    name = re.sub(r"[^a-z ]", "", name.lower().replace("-", " "))
    name = " ".join(w for w in name.split() if w not in ("jr", "sr", "ii", "iii", "iv", "v"))
    return FP_NAME_ALIASES.get(name, name)


def league_owners(teams):
    """Map (normalized name, position) -> list of (nfl team, owner handle) for rostered players."""
    players = api("players/nfl")
    owners = {}
    for r in api(f"league/{LEAGUE_ID}/rosters"):
        handle = teams[r["roster_id"]]["handle"]
        for pid in r.get("players") or []:
            p = players.get(pid)
            if p and p.get("position") in FP_POSITIONS:
                key = (_norm(p.get("full_name") or ""), p["position"])
                owners.setdefault(key, []).append((p.get("team") or "FA", handle))
    return owners


def owner_of(owners, player, pos, team):
    matches = owners.get((_norm(player), pos), [])
    if len(matches) > 1:
        matches = [m for m in matches if m[0] == team] or matches
    return matches[0][1] if matches else "FA"


def build_player_stats(wb, fp):
    ws = wb.create_sheet(f"Player Stats ({SEASON})")
    ws["B2"], ws["C2"] = "Data", "FantasyPros (Half PPR)"
    ws["B2"].font = TITLE_FONT
    col = 2
    for pos in FP_POSITIONS:
        block = fp[pos]
        ws.cell(4, col, f"{pos} Fantasy Scoring (0.5 PPR)").font = TITLE_FONT
        head = [("", "Week"), ("", "Team"), ("", "Player")] + block["cols"]
        for j, (grp, name) in enumerate(head):
            g = ws.cell(5, col + j, grp or None)
            g.font = Font(bold=True, size=9, color="6B7280")
        header(ws, 6, col, [n for _, n in head])
        for i, row in enumerate(sorted(block["rows"], key=lambda r: (r["week"], -(r["stats"].get(("MISC", "FPTS")) or 0)))):
            vals = [row["week"], f"({row['team']})", row["player"]] + [row["stats"].get(c) for c in block["cols"]]
            for j, v in enumerate(vals):
                ws.cell(7 + i, col + j, v)
        ws.column_dimensions[get_column_letter(col + 2)].width = 22
        col += len(head) + 1
    ws.freeze_panes = "A7"


def player_totals(fp, owners, pos):
    """Season totals per player at one position, tagged with their league owner."""
    totals = {}
    for row in fp[pos]["rows"]:
        t = totals.setdefault(row["fp_id"], {"player": row["player"], "team": row["team"], "g": 0, "fpts": 0.0, "wk": 0})
        t["g"] += row["stats"].get(("MISC", "G")) or 0
        t["fpts"] += row["stats"].get(("MISC", "FPTS")) or 0
        if row["week"] >= t["wk"]:  # keep the latest team after trades
            t["team"], t["wk"] = row["team"], row["week"]
    for t in totals.values():
        t["fpts"] = round(t["fpts"], 2)
        t["ppg"] = round(t["fpts"] / t["g"], 2) if t["g"] else 0
        t["owner"] = owner_of(owners, t["player"], pos, t["team"])
    return totals


def build_player_rankings(wb, fp, owners, n_weeks):
    ws = wb.create_sheet(f"Player Rankings ({SEASON})")
    title(ws, f"{SEASON} Fantasy Ranks – Half PPR, weeks 1-{n_weeks} (owner = Go Big or Go Home roster)")
    col = 2
    for pos in FP_POSITIONS:
        totals = player_totals(fp, owners, pos)
        for key, label in (("fpts", "FPTS"), ("ppg", "FPTS / G")):
            ranked = sorted(totals.values(), key=lambda t: -t[key])
            ws.cell(3, col, f"{pos} Ranking – {'Total FPTS' if key == 'fpts' else 'FPTS / G'}").font = Font(bold=True)
            header(ws, 4, col, ["Rank", "Player", "Team", "G", label, "Owner"])
            for i, (t, rk) in enumerate(zip(ranked, comp_rank([t[key] for t in ranked]))):
                for j, v in enumerate([rk, t["player"], t["team"], t["g"], t[key], t["owner"]]):
                    c = ws.cell(5 + i, col + j, v)
                    if t["owner"] == "@MGardner6":
                        c.fill = ME_FILL
                    elif j == 5 and v == "FA":
                        c.font = WIN_FONT
            for j, wdt in enumerate((6, 22, 6, 5, 9, 16)):
                ws.column_dimensions[get_column_letter(col + j)].width = wdt
            col += 7
    ws.freeze_panes = "A5"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--weeks", type=int)
    ap.add_argument("--skip-fantasypros", action="store_true")
    args = ap.parse_args()
    league, weeks, reg_weeks, order, teams = load_season(args.weeks)
    wb = Workbook()
    build_heatmap(wb, weeks, reg_weeks, order)
    build_weekly(wb, weeks, reg_weeks, order, teams)
    build_pfpa(wb, weeks, reg_weeks, order)
    build_league_data(wb, weeks, order, teams)
    build_managers(wb, order)
    build_payments(wb, order)
    if not args.skip_fantasypros:
        fp = load_fantasypros(weeks)
        owners = league_owners(teams)
        build_player_rankings(wb, fp, owners, len(weeks))
        build_player_stats(wb, fp)
        rostered = sum(len(v) for v in owners.values())
        matched = {(_norm(r["player"]), pos) for pos in FP_POSITIONS for r in fp[pos]["rows"]}
        missing = sum(len(v) for k, v in owners.items() if k not in matched)
        counts = ", ".join(f"{p} {len(fp[p]['rows'])}" for p in FP_POSITIONS)
        print(f"FantasyPros: {counts} player-weeks; "
              f"{rostered - missing}/{rostered} rostered QB/RB/WR/TE matched to a FantasyPros stat line")
    wb.save(OUT)
    print(f"Saved {OUT} (weeks {weeks[0]}-{weeks[-1]})")
    for t in order:
        print(f"{t['place']:>2}. {t['handle']:<16} {t['w']}-{t['l']}  PF {t['pf']:>7.2f}  "
              f"PA {t['pa']:>7.2f}  all-play {t['ap_w']}-{t['ap_l']}  luck {t['luck']:+d}")


if __name__ == "__main__":
    main()
